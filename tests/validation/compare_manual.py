"""Compare the program against the supervisor's manual xlsx reference.

VALIDATION A -- formula check. CAN FAIL (exit 1), tolerance 1e-9.
    Feeds the xlsx's OWN counts and per-element Q.Omega/sigma into
    calculate_uncertainty_v9 and compares against the xlsx uncertainties,
    per layer and per element. This isolates the formula from the ROI
    question entirely: same inputs in, same numbers expected out.

VALIDATION B -- counts comparison. REPORT ONLY, never fails.
    Program A_A for the rois.json ROIs beside the xlsx counts. While
    rois.json holds placeholder ROIs these deltas are expected to be large;
    that is the signal, not a failure.

Layer mapping: xlsx L1-L5 == xnra layers 2-6 (the ZrOx film).
H is excluded: it has no column in this RBS element spectrum.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import _paths  # noqa: E402

_paths.add_repo_to_path()

from xnra_element_spectra_v1 import (  # noqa: E402
    calculate_element_counts_from_spectrum,
    parse_element_spectrum,
)
from xnra_uncertainty_v12 import calculate_uncertainty_v9, parse_xnra_full  # noqa: E402

TOL = 1e-9

# xlsx cell layout of the per-element uncertainty blocks:
# (value column, first row, counts column); rows are stat / charge / total / xsec
XLSX_BLOCKS = {
    "H": ("B", 20, "C"),
    "C": ("F", 20, "G"),
    "O": ("B", 24, "C"),
    "Zr": ("F", 24, "G"),
    "Hf": ("B", 28, "C"),
}
ELEMENTS = ["C", "O", "Zr", "Hf"]
XLSX_TO_XNRA_LAYER = {1: 2, 2: 3, 3: 4, 4: 5, 5: 6}


def read_xlsx():
    import openpyxl

    wb = openpyxl.load_workbook(str(_paths.XLSX), data_only=True)
    ws = wb[_paths.XLSX_SHEET]
    out = {}
    for element, (col, row, ccol) in XLSX_BLOCKS.items():
        out[element] = {
            "stat_rel": ws[f"{col}{row}"].value,
            "counts": ws[f"{ccol}{row}"].value,
            "qomega": ws[f"{col}{row + 1}"].value,
            "total_rel": ws[f"{col}{row + 2}"].value,
            "xsec": ws[f"{col}{row + 3}"].value,
        }
    return out


def validation_a(data, xlsx):
    print("=" * 100)
    print("VALIDATION A -- formula check (xlsx counts + xlsx Q.Omega/sigma -> program)")
    print(f"tolerance {TOL:g}; CAN FAIL")
    print("=" * 100)
    header = (
        f"{'xlsx':<5}{'xnra':<5}{'El':<4}{'Conc':>8}{'A_A':>10}"
        f"{'rel xlsx':>14}{'rel prog':>14}{'d rel':>11}"
        f"{'abs xlsx':>14}{'abs prog':>14}{'d abs':>11}  "
    )
    print(header)
    print("-" * 100)

    failures = 0
    for xlsx_layer, xnra_layer in sorted(XLSX_TO_XNRA_LAYER.items()):
        layer = next(l for l in data["layers"] if l["layer_num"] == xnra_layer)
        for element in ELEMENTS:
            ref = xlsx[element]
            conc = layer["concentrations"].get(element, 0)
            if conc <= 0:
                continue

            unc = calculate_uncertainty_v9(
                conc,
                ref["counts"],
                ref["qomega"],
                ref["xsec"],
                adjustment_method="none",
            )
            prog_rel = float(unc["sigma_total_rel"])
            prog_abs = float(unc["sigma_total"])
            ref_rel = ref["total_rel"]
            ref_abs = ref_rel * conc

            d_rel = abs(prog_rel - ref_rel)
            d_abs = abs(prog_abs - ref_abs)
            bad = d_rel > TOL or d_abs > TOL
            failures += bad

            print(
                f"L{xlsx_layer:<4}L{xnra_layer:<4}{element:<4}{conc:>8.4f}{ref['counts']:>10d}"
                f"{ref_rel:>14.10f}{prog_rel:>14.10f}{d_rel:>11.2e}"
                f"{ref_abs:>14.10f}{prog_abs:>14.10f}{d_abs:>11.2e}  "
                f"{'FAIL' if bad else 'ok'}"
            )

    print("-" * 100)
    if failures:
        print(f"VALIDATION A: FAIL -- {failures} row(s) exceed {TOL:g}")
    else:
        print(f"VALIDATION A: PASS -- all rows within {TOL:g}")
    return failures


def validation_b(data, spectrum, cfg, xlsx):
    print()
    print("=" * 100)
    print("VALIDATION B -- program A_A vs xlsx counts (REPORT ONLY, never fails)")
    print(f"roi_source = {cfg.get('roi_source')} -- {cfg.get('_comment', '')}")
    print("=" * 100)
    print(
        f"{'El':<4}{'ROI':>14}{'A_A prog':>14}{'counts xlsx':>14}"
        f"{'delta':>14}{'ratio':>9}{'d stat_rel':>13}"
    )
    print("-" * 100)

    for entry in cfg["elements"]:
        element = entry["element"]
        if element not in xlsx:
            continue
        roi_start, roi_end = entry["channels"]
        counts, _fraction, _exp = calculate_element_counts_from_spectrum(
            spectrum, data["raw_counts"], element, roi_start, roi_end, use_isotope=False
        )
        ref_counts = xlsx[element]["counts"]
        prog_stat = counts ** -0.5 if counts > 0 else float("nan")
        ref_stat = xlsx[element]["stat_rel"]
        print(
            f"{element:<4}{f'{roi_start}-{roi_end}':>14}{counts:>14d}{ref_counts:>14d}"
            f"{counts - ref_counts:>+14d}{counts / ref_counts:>9.4f}"
            f"{prog_stat - ref_stat:>+13.2e}"
        )

    print("-" * 100)
    print("A_A = integrated SIMULATED counts in the ROI")
    print("     (xnra_element_spectra_v1.py:404; experimental sum is diagnostic only).")
    print("Deltas are expected while rois.json holds placeholder ROIs.")


def main():
    _paths.require(_paths.XNRA, _paths.METADATA, _paths.XLSX, _paths.ROIS_JSON)

    cfg = _paths.load_rois()
    data = parse_xnra_full(str(_paths.XNRA))
    spectrum = parse_element_spectrum(str(_paths.METADATA))
    xlsx = read_xlsx()

    failures = validation_a(data, xlsx)
    validation_b(data, spectrum, cfg, xlsx)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
